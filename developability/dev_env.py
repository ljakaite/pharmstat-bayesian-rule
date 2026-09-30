"""
Developability stage-gate environment (c_SREP_10.md §4 option 1, c_SREP_11.md
§3), using the real Jain et al. (2017, PNAS 114(5):944-949) dataset: 137
clinical-stage antibodies, 12 biophysical developability assays each.

Files (downloaded by the user, PNAS blocks automated fetch):
  pnas.1616408114.sd01.xlsx -- antibody list + clinical status
  pnas.1616408114.sd02.xlsx -- VH/VL sequences
  pnas.1616408114.sd03.xlsx -- the 12-assay results table (this module's data)

**Checked and found NOT present in this exact file, flagged rather than
assumed:** the dual-condition (PBS pH 7.4 vs. His/Arg pH 6.0) self-association
measurement mentioned in earlier research (c_SREP_9/10/11.md) is not in this
dataset -- there is only ONE AC-SINS column, single condition. No genuine
switching regime exists here either, same conclusion as for GB1. R is
dropped for this environment too, per the same "don't invent dynamics the
data doesn't have" principle already applied to GB1's regime.

Composite "true developability" score: mean of 12 assays' z-scores across the
137-antibody population, each SIGN-FLIPPED so higher z = better, per standard
interpretation in this literature (checked, not guessed):
  higher-is-better:  HEK Titer, Fab Tm, SGAC-SINS AS100
  lower-is-better:   HIC/SMAC/CIC retention time, accelerated-stability slope,
                      PSR, AC-SINS, CSI-BLI, ELISA, BVP ELISA (all
                      hydrophobicity/self-interaction/polyspecificity assays)
"""
import statistics as st
from pathlib import Path

import openpyxl

DATA_DIR = Path(__file__).parent / "data" / "pnas_xlsx"
SD01 = DATA_DIR / "pnas.1616408114.sd01.xlsx"
SD03 = DATA_DIR / "pnas.1616408114.sd03.xlsx"

ASSAY_COLUMNS = [
    "HEK Titer (mg/L)", "Fab Tm by DSF (°C)", "SGAC-SINS AS100 ((NH4)2SO4 mM)",
    "HIC Retention Time (Min)a", "SMAC Retention Time (Min)a",
    "Slope for Accelerated Stability", "Poly-Specificity Reagent (PSR) SMP Score (0-1)",
    "Affinity-Capture Self-Interaction Nanoparticle Spectroscopy (AC-SINS) ∆λmax (nm) Average",
    "CIC Retention Time (Min)", "CSI-BLI Delta Response (nm)", "ELISA", "BVP ELISA",
]
HIGHER_IS_BETTER = {"HEK Titer (mg/L)", "Fab Tm by DSF (°C)", "SGAC-SINS AS100 ((NH4)2SO4 mM)"}

# Three real assay-subset panels standing in for Grok's "competing predictors"
# (sequence-only / structure-based / thermodynamic) -- adapted, not identical:
# this session has no structure-prediction tooling available, so the three
# competing models are three real, distinct EARLY-STAGE assay subsets used to
# predict the FULL 12-assay composite, rather than three literally different
# modelling paradigms. Flagged as an honest adaptation, matching the same
# discipline applied to the RDM-analogue caveat in c_SREP_9.md.
PANELS = {
    "expression_stability": ["HEK Titer (mg/L)", "Fab Tm by DSF (°C)"],
    "aggregation": ["HIC Retention Time (Min)a", "SMAC Retention Time (Min)a",
                     "SGAC-SINS AS100 ((NH4)2SO4 mM)",
                     "Affinity-Capture Self-Interaction Nanoparticle Spectroscopy (AC-SINS) ∆λmax (nm) Average"],
    "specificity": ["Poly-Specificity Reagent (PSR) SMP Score (0-1)", "CIC Retention Time (Min)",
                     "ELISA", "BVP ELISA"],
}


def load_assays():
    wb = openpyxl.load_workbook(SD03, data_only=True)
    ws = wb["Results-12-assays"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    col_idx = {h: i for i, h in enumerate(header) if h}
    data = {}
    for row in rows[1:]:
        name = row[0]
        if not name:
            continue
        vals = {}
        ok = True
        for col in ASSAY_COLUMNS:
            v = row[col_idx[col]]
            if not isinstance(v, (int, float)):
                ok = False
                break
            vals[col] = v
        if ok:
            data[name] = vals
    return data


def load_clinical_status():
    wb = openpyxl.load_workbook(SD01, data_only=True)
    ws = wb["Antibody-list"]
    rows = list(ws.iter_rows(values_only=True))
    return {r[0]: r[4] for r in rows[1:] if r[0]}


def compute_composite(data: dict):
    """z-score each assay across the population, sign-flip so higher=better,
    average across all 12 -> one real, checkable "true developability" scalar
    per antibody (the hidden ground truth the agent doesn't see directly)."""
    zscored = {}
    for col in ASSAY_COLUMNS:
        vals = [data[name][col] for name in data]
        mu, sd = st.mean(vals), st.pstdev(vals)
        sign = 1.0 if col in HIGHER_IS_BETTER else -1.0
        for name in data:
            z = sign * (data[name][col] - mu) / sd if sd > 0 else 0.0
            zscored.setdefault(name, {})[col] = z
    composite = {name: sum(zscored[name].values()) / len(ASSAY_COLUMNS) for name in data}
    return composite, zscored


def fit_panel_models(data: dict, composite: dict):
    """Fit one linear model per panel: predict the full composite from that
    panel's assay subset alone, using the real population -- three genuine,
    checkable predictors, not hand-picked constants."""
    from sklearn.linear_model import LinearRegression
    names = list(data.keys())
    models = {}
    errors = {}
    for panel, cols in PANELS.items():
        X = [[data[n][c] for c in cols] for n in names]
        y = [composite[n] for n in names]
        reg = LinearRegression().fit(X, y)
        models[panel] = (reg, cols)
        preds = reg.predict(X)
        errors[panel] = sum(abs(p - t) for p, t in zip(preds, y)) / len(y)
    return models, errors


def panel_predict(models: dict, name: str, data: dict):
    preds = {}
    for panel, (reg, cols) in models.items():
        x = [[data[name][c] for c in cols]]
        preds[panel] = float(reg.predict(x)[0])
    return preds
